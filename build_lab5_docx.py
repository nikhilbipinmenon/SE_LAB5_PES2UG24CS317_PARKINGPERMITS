import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

# Title
p_title = doc.add_paragraph()
r_uni = p_title.add_run("PES UNIVERSITY\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
r_uni.bold = True
r_uni.font.color.rgb = RGBColor(0, 82, 204)
r_sub = p_title.add_run("Software Engineering Lab (UE23CS252A) — Academic Year 2026-2027\n")
r_sub.font.size = Pt(10)
r_sub.font.color.rgb = RGBColor(94, 108, 132)

p_h = doc.add_heading("Lab 5 Deliverables: Static Code Analysis", level=1)
p_meta = doc.add_paragraph()
p_meta.add_run("Topic: Campus Parking Permit Manager (parking_system.py)\nStudent: Nikhil B Menon | SRN: PES2UG24CS317\n").italic = True

# Section 1
doc.add_heading("1. Executive Summary & Verification Metrics", level=2)
table = doc.add_table(rows=4, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Tool", "Initial State (PRE / Before Fixes)", "Final Cleaned State (POST / After Fixes)", "Status"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True

rows_data = [
    ["Pylint", "Score: 6.32 / 10 (Mutable args, bare excepts, missing docstrings, global state)", "Score: 10.00 / 10 (+3.68 improvement, zero defects)", "PASSED (10/10)"],
    ["Bandit", "2 Security Issues (B311: Insecure PRNG, B110: Try-except-pass)", "No issues identified (0 High, 0 Medium, 0 Low)", "PASSED (Clean)"],
    ["Flake8", "8 Style Errors (E302 spacing, E501 line length >79 chars, E722 bare excepts)", "0 Errors / 0 Warnings (100% PEP 8 compliant)", "PASSED (Clean)"]
]
for r_idx, r_data in enumerate(rows_data):
    for c_idx, val in enumerate(r_data):
        table.cell(r_idx + 1, c_idx).text = val

doc.add_paragraph()

# Section 2: Terminal Screenshots
doc.add_heading("2. Terminal Verification Screenshots (PRE & POST Runs)", level=2)
doc.add_paragraph("Figure 1: PRE-Analysis Terminal Run (Before Fixes) — Demonstrating initial defects (Bandit 2 issues, Flake8 violations):")
doc.add_picture("/Users/nikhilbmenon/Sem 5 Labs/SE Lab/L5_10_parking_permits/screenshot_before_analysis.png", width=Inches(6.2))

doc.add_paragraph("Figure 2: POST-Analysis Terminal Run (After Fixes) — Demonstrating verified clean code (Pylint 10.00/10, Bandit clean):")
doc.add_picture("/Users/nikhilbmenon/Sem 5 Labs/SE Lab/L5_10_parking_permits/screenshot_after_analysis.png", width=Inches(6.2))

# Section 3: Issues Table
doc.add_page_break()
doc.add_heading("3. Issues Documentation Table (From issues.md)", level=2)
t_issues = doc.add_table(rows=10, cols=5)
t_issues.alignment = WD_TABLE_ALIGNMENT.CENTER
i_headers = ["Issue Defect", "Type & Tool", "Line", "Description", "Fix Approach"]
for i, h in enumerate(i_headers):
    t_issues.cell(0, i).text = h
    t_issues.cell(0, i).paragraphs[0].runs[0].bold = True

issues_data = [
    ["Insecure PRNG", "Security (Bandit B311)", "11", "Using random.randint() for permit IDs is cryptographically predictable.", "Replaced with secure secrets.randbelow(90000) + 10000."],
    ["Bare Except", "Security (Bandit B110 / Pylint W0702 / Flake8 E722)", "22", "Bare except: pass silently conceals critical runtime exceptions.", "Replaced with atomic permits.pop() and specific handlers."],
    ["Mutable Default", "Bug (Pylint W0102)", "9", "zone_records=[] retains list mutations across multiple calls.", "Default set to None; initialized inside method."],
    ["Unclosed File", "Resource Leak (Pylint R1732 / W1514)", "42, 47", "Calling open() without context manager risks leaks; missing UTF-8.", "Refactored to with open(file, mode, encoding='utf-8') as f:."],
    ["Global State", "Code Smell (Pylint W0603)", "40", "Function mutated global permits, breaking encapsulation.", "Updated in-place via permits.clear() and permits.update()."],
    ["Naming Style", "Style (Pylint C0103)", "51", "Function setExpiry used camelCase instead of PEP 8 snake_case.", "Renamed to set_expiry with backward-compatible alias."],
    ["Missing Docs", "Convention (Pylint C0114 / C0116)", "Module", "Module and 7 functions lacked docstrings explaining behavior & params.", "Added comprehensive PEP 257 compliant docstrings."],
    ["PEP 8 Format", "Style (Flake8 E302 / E501)", "Various", "Lines exceeded 79 characters; single blank line instead of two.", "Formatted spacing to PEP 8 standards and wrapped long lines."],
    ["Logic Bugs", "Domain Logic Bug", "56–60", "Accepted empty plates and invalid zones; expiry crashed on date format.", "Added plate sanitization, zone membership check, and date parsing."]
]
for r_idx, r_data in enumerate(issues_data):
    for c_idx, val in enumerate(r_data):
        t_issues.cell(r_idx + 1, c_idx).text = val

doc.add_paragraph()

# Section 4: Reflection
doc.add_heading("4. Answers to Reflection Questions (From reflection.md)", level=2)

questions = [
    ("Question 1: Tool Comparison & Effectiveness", "Bandit and Pylint uncovered the most critical and actionable issues. Bandit flagged high-impact security vulnerabilities: B311 (insecure pseudo-random number generation) and B110 (bare except: pass). Pylint detected an insidious Python logic bug: the mutable default argument (zone_records=[]). Flake8 was essential for enforcing PEP 8 style formatting."),
    ("Question 2: Contribution to Quality & Security", "Static code analysis enables Shift-Left Quality and Security. It acts as an automated first line of defense, discovering latent resource leaks, type errors, and cryptographic weaknesses before runtime. Furthermore, enforcing strict conventions (PEP 8) ensures readability and maintainability across teams."),
    ("Question 3: Challenges and Limitations", "Static analysis tools are context-insensitive and cannot reason about business domain semantics. While the tools flagged line lengths and library anti-patterns, they completely missed business logic bugs (e.g. accepting invalid parking zone 'Z' or empty plates). Thus, static analysis must always be paired with functional unit testing."),
    ("Question 4: CI/CD Pipeline Integration", "In a CI/CD pipeline (e.g. GitHub Actions), static analysis tools run on every pull request as automated quality gates. Merges are blocked if Bandit flags High/Medium vulnerabilities or if the Pylint score falls below 9.0/10. Pre-commit hooks run Flake8 locally before committing, and findings are exported in SARIF format to the GitHub Security tab.")
]
for q_title, q_ans in questions:
    p_q = doc.add_paragraph()
    p_q.add_run(q_title + "\n").bold = True
    p_q.add_run(q_ans)

docx_path = "/Users/nikhilbmenon/Sem 5 Labs/SE Lab/L5_10_parking_permits/Lab5_Static_Code_Analysis_Report.docx"
doc.save(docx_path)
import shutil
shutil.copy2(docx_path, "/Users/nikhilbmenon/Sem 5 Labs/SE Lab/Lab5_Static_Code_Analysis_Report.docx")
print("Saved DOCX deliverable at:", docx_path)
